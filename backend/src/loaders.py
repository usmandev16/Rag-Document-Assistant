import logging
import os
import re
from io import BytesIO

import pandas as pd
from docx import Document
from pypdf import PdfReader

from src.chunking import chunk_text

_DATE_NAME_RE = re.compile(r"date|time", re.IGNORECASE)

# A PDF page with fewer extractable characters than this is OCR'd. 200, not
# 50: some PDFs yield only a ~57-char running header per page while the body
# is unextractable, and a real text page has far more than 200.
OCR_MIN_CHARS = int(os.environ.get("OCR_MIN_CHARS", "200"))
OCR_DPI = 300

# Own handler so the per-document OCR count shows in the Celery worker log.
log = logging.getLogger("rag.ingest")
if not log.handlers:
    _handler = logging.StreamHandler()
    _handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(name)s: %(message)s"))
    log.addHandler(_handler)
    log.setLevel(logging.INFO)
    log.propagate = False


def _fix_excel_serial_dates(df: pd.DataFrame) -> pd.DataFrame:
    """openpyxl only parses a cell as a date when the source file's own
    number-format says so; a date column stored as plain numbers survives as
    a raw Excel serial (e.g. 46204) instead of a real date. Reinterpret any
    numeric column whose name looks date-like as an Excel serial date."""
    for col in df.columns:
        if not _DATE_NAME_RE.search(str(col)):
            continue
        if pd.api.types.is_numeric_dtype(df[col]):
            df[col] = pd.to_datetime(
                df[col], unit="D", origin="1899-12-30", errors="coerce"
            ).dt.strftime("%Y-%m-%d")
    return df


ROWS_PER_CHUNK = 25


def _dataframe_chunks(df: pd.DataFrame, label: str = "") -> list[str]:
    """Split a table into chunks of ROWS_PER_CHUNK whole rows, each starting
    with the header row ("Col A | Col B | ...") so every chunk is
    self-describing. Rows are never cut, so an ID inside one stays intact.
    Header-once instead of "Col: value" per cell keeps chunks much smaller
    in tokens (Groq free tier: 8k tokens/minute).
    """
    df = _fix_excel_serial_dates(df.copy()).fillna("")
    header = " | ".join(str(c) for c in df.columns)
    rows = [" | ".join(str(v) for v in row) for row in df.itertuples(index=False)]
    prefix = f"{label}\n" if label else ""
    return [
        f"{prefix}{header}\n" + "\n".join(rows[i:i + ROWS_PER_CHUNK])
        for i in range(0, len(rows), ROWS_PER_CHUNK)
    ]


def _ocr_page(file_bytes: bytes, page_index: int) -> str:
    """Render one PDF page at OCR_DPI and read it with Tesseract.

    imported lazily so the API and local dev don't need Tesseract;
    only the Celery worker image installs it (see Dockerfile). Reopens the
    PDF per OCR'd page, which is cheap next to the OCR itself."""
    import pypdfium2 as pdfium
    import pytesseract

    page = pdfium.PdfDocument(file_bytes)[page_index]
    image = page.render(scale=OCR_DPI / 72).to_pil()
    return pytesseract.image_to_string(image)


def _pdf_chunks(file_name: str, file_bytes: bytes) -> list[tuple[str, dict]]:
    """Chunk a PDF page by page so every chunk keeps its page number. Pages
    where pypdf finds almost no text (scanned, or text drawn as shapes —
    often only a running header survives) are OCR'd instead."""
    reader = PdfReader(BytesIO(file_bytes))
    chunks: list[tuple[str, dict]] = []
    ocr_pages = 0
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        ocr = len(text.strip()) < OCR_MIN_CHARS
        if ocr:
            text = _ocr_page(file_bytes, number - 1)
            ocr_pages += 1
        chunks += [(piece, {"page": number, "ocr": ocr}) for piece in chunk_text(text, chunk_size=800, overlap=150)]
    log.info("%s: %d of %d pages used OCR", file_name, ocr_pages, len(reader.pages))
    return chunks


def extract_chunks(file_name: str, file_bytes: bytes) -> list[tuple[str, dict]]:
    """File -> (chunk text, extra metadata) pairs. PDFs are chunked per page
    (metadata: page, ocr); spreadsheets by rows with the header repeated;
    everything else through the paragraph-packing text chunker."""
    name = file_name.lower()
    if name.endswith(".pdf"):
        return _pdf_chunks(file_name, file_bytes)
    if name.endswith(".xlsx"):
        sheets = pd.read_excel(BytesIO(file_bytes), sheet_name=None)
        return [(c, {}) for sheet, df in sheets.items() for c in _dataframe_chunks(df, f"# Sheet: {sheet}")]
    if name.endswith(".csv"):
        return [(c, {}) for c in _dataframe_chunks(pd.read_csv(BytesIO(file_bytes)))]
    return [(c, {}) for c in chunk_text(extract_text(file_name, file_bytes), chunk_size=800, overlap=150)]


def extract_text(file_name: str, file_bytes: bytes) -> str:
    name = file_name.lower()
    if name.endswith(".docx"):
        document = Document(BytesIO(file_bytes))
        return "\n\n".join(p.text for p in document.paragraphs if p.text)
    return file_bytes.decode("utf-8", errors="ignore")
